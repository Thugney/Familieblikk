"""Preserve and stage a private workbook for a future household import.

No database writes, formula execution, external requests, inferred owner identity,
or modification of the source workbook. Use the bundled Python/openpyxl runtime.
"""
from __future__ import annotations
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re
import uuid
import warnings
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PRIVATE = ROOT / 'private-data' / 'family-budget'
SOURCE = PRIVATE / 'source' / 'family-budget-expandable.xlsx'
IMPORT_ID = 'family-budget-expandable-v1'
KINDS = {'Income': 'income', 'Expense': 'expense', 'Debt': 'debt_payment',
         'Savings': 'savings_deposit', 'Annual': 'annual_fund_payment',
         'Sent abroad': 'remittance'}
NAMESPACE = uuid.UUID('d2c1bb30-7956-4b51-a42f-0f79aa7cbb11')

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def serial(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value

def save(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=serial) + '\n', encoding='utf-8')
    temp.replace(path)

def stable_id(sheet, row):
    # A workbook revision is NOT part of the key: reimporting an edited copy must
    # not append every old row again. The future DB must also scope it to household.
    return str(uuid.uuid5(NAMESPACE, f'{IMPORT_ID}:{sheet}:{row}'))

def money_ore(value):
    if value is None or value == '' or isinstance(value, bool):
        return None
    if not isinstance(value, (int, float, Decimal)):
        raise ValueError('not_numeric')
    d = Decimal(str(value))
    if not d.is_finite():
        raise ValueError('not_finite')
    if d != d.quantize(Decimal('.01')):
        raise ValueError('more_than_two_decimals')
    return int(d * 100)

def decimal_value(value):
    if value is None or value == '' or isinstance(value, bool):
        return None
    if not isinstance(value, (int, float, Decimal)):
        raise ValueError('not_numeric')
    d = Decimal(str(value))
    if not d.is_finite():
        raise ValueError('not_finite')
    return d

def calendar_date(value):
    if isinstance(value, datetime):
        if value.time().isoformat() != '00:00:00':
            raise ValueError('non_midnight_timestamp')
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None

def meaningful(value):
    return value is not None and value != ''

def staging_eligibility(issues):
    # No local manifest can authenticate a user or prove an available application.
    # Real application import must recheck its session, tenant and review state.
    review_count = sum(i['severity'] == 'review' for i in issues)
    return {'financial_review_issues': review_count,
            'all_financial_records_resolved': review_count == 0,
            'safe_to_apply_as_confirmed_ledger': False,
            'owner_verified': False, 'application_available': False,
            'blockers': (['unresolved_source_data'] if review_count else []) +
                        ['verified_owner_session_required', 'live_application_required']}

class Extractor:
    def __init__(self, source=SOURCE):
        self.source = Path(source)
        with warnings.catch_warnings():
            # Reading does not save, remove, or rewrite unsupported Excel extensions.
            warnings.simplefilter('ignore', UserWarning)
            self.book = openpyxl.load_workbook(self.source, read_only=False, data_only=False, keep_links=False)
            self.cache = openpyxl.load_workbook(self.source, read_only=False, data_only=True, keep_links=False)
        self.issues = []
        self.coverage = defaultdict(list)
        self.controls = []

    def issue(self, record, code, field=None, severity='review', detail=None):
        item = {'record_id': record.get('source_id'), 'source': record.get('source'),
                'code': code, 'field': field, 'severity': severity}
        if detail:
            item['detail'] = detail
        self.issues.append(item)
        record.setdefault('issue_codes', []).append(code)

    def row(self, sheet, number, fields):
        cells = {key: self.book[sheet].cell(number, col) for key, col in fields.items()}
        raw = {key: serial(c.value) for key, c in cells.items()}
        record = {'source_id': stable_id(sheet, number), 'source': {'sheet': sheet, 'row': number},
                  'raw': raw, 'cells': {key: c.coordinate for key, c in cells.items()},
                  'issue_codes': []}
        for key, c in cells.items():
            if c.data_type == 'f':
                record.setdefault('formulas', {})[key] = c.value
                record.setdefault('cached_values', {})[key] = serial(self.cache[sheet][c.coordinate].value)
            if isinstance(c.value, str) and '\ufffd' in c.value:
                self.issue(record, 'source_text_contains_replacement_character', key, 'warning')
        return record

    def entered(self, sheet, row, columns):
        return any(meaningful(self.book[sheet].cell(row, c).value)
                   and self.book[sheet].cell(row, c).data_type != 'f' for c in columns)

    def amount(self, rec, field, target, allow_negative=False):
        if field in rec.get('formulas', {}):
            self.issue(rec, 'input_is_formula_not_confirmed_amount', field)
            rec[target] = None
            return
        try:
            rec[target] = money_ore(rec['raw'].get(field))
            if rec[target] is None:
                self.issue(rec, 'missing_amount', field)
            elif rec[target] < 0 and not allow_negative:
                self.issue(rec, 'unexpected_negative_amount', field)
        except (ValueError, InvalidOperation) as e:
            rec[target] = None
            self.issue(rec, 'invalid_amount', field, detail=str(e))

    def day(self, rec, field, target, required=True):
        try:
            value = self.book[rec['source']['sheet']][rec['cells'][field]].value
            rec[target] = calendar_date(value)
            if rec[target] is None and (required or meaningful(value)):
                self.issue(rec, 'missing_or_ambiguous_date', field)
        except ValueError as e:
            rec[target] = None
            self.issue(rec, 'ambiguous_date', field, detail=str(e))

    def rows(self, sheet, start, end, fields, skip_template=False):
        result = []
        for r in range(start, min(end, self.book[sheet].max_row) + 1):
            if not self.entered(sheet, r, fields.values()):
                continue
            rec = self.row(sheet, r, fields)
            if skip_template:
                name = rec['raw'].get('name')
                if isinstance(name, str) and name.startswith('---'):
                    self.controls.append({'sheet': sheet, 'row': r, 'reason': 'section_heading'})
                    continue
                if not meaningful(name) and all(not meaningful(rec['raw'].get(k)) or rec['raw'].get(k) == 0
                                                for k in ['monthly_budget', 'note']):
                    self.controls.append({'sheet': sheet, 'row': r, 'reason': 'unused_typed_template_row'})
                    continue
            result.append(rec)
            self.coverage[sheet].append(r)
        return result

    def snapshot(self):
        sheets = []
        for ws in self.book.worksheets:
            cells = []
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None:
                        continue
                    item = {'cell': c.coordinate, 'type': c.data_type, 'value': serial(c.value),
                            'number_format': c.number_format}
                    if c.data_type == 'f':
                        item['cached_value'] = serial(self.cache[ws.title][c.coordinate].value)
                    cells.append(item)
            sheets.append({'name': ws.title, 'state': ws.sheet_state, 'cells': cells})
        return {'schema_version': 1, 'source_sha256': digest(self.source),
                'formulas_executed': False, 'sheets': sheets,
                'defined_names': {k: n.attr_text for k, n in self.book.defined_names.items()}}

    def extract(self):
        settings = self.row('Settings', 3, {'household_name': 2})
        settings['raw'] = {'household_name': self.book['Settings']['B3'].value,
                           'base_currency': self.book['Settings']['B4'].value,
                           'budget_year': self.book['Settings']['B5'].value,
                           'selected_month': self.book['Dashboard']['C3'].value,
                           'debt_strategy': self.book['Debts']['B5'].value,
                           'debt_extra_monthly': self.book['Debts']['B6'].value}
        settings['source_cells'] = {'household_name': 'Settings!B3', 'base_currency': 'Settings!B4',
                                   'budget_year': 'Settings!B5', 'selected_month': 'Dashboard!C3',
                                   'debt_strategy': 'Debts!B5', 'debt_extra_monthly': 'Debts!B6'}
        currency = settings['raw']['base_currency']
        if currency != 'NOK':
            self.issue(settings, 'source_base_currency_differs_from_approved_nok', 'base_currency')

        categories = self.rows('Categories', 6, self.book['Categories'].max_row,
                               {'name': 1, 'kind': 2, 'monthly_budget': 3, 'note': 4}, skip_template=True)
        by_name = defaultdict(list)
        for c in categories:
            name = c['raw']['name']
            c['accounting_class'] = KINDS.get(c['raw']['kind'])
            if not meaningful(name):
                self.issue(c, 'missing_category_name', 'name')
            else:
                by_name[name].append(c)
            if not c['accounting_class']:
                self.issue(c, 'unknown_category_kind', 'kind')
            if 'monthly_budget' in c.get('formulas', {}):
                c['monthly_budget_ore'] = None
                c['budget_mode'] = 'derived_from_linked_module_plan'
            else:
                self.amount(c, 'monthly_budget', 'monthly_budget_ore')
                c['budget_mode'] = 'source_monthly_limit'
            c['effective_month'] = None
        for rows in by_name.values():
            if len(rows) > 1:
                for r in rows:
                    self.issue(r, 'duplicate_category_name', 'name')

        debts = self.rows('Debts', 11, self.book['Debts'].max_row - 1,
                          {'name': 1, 'type': 2, 'balance': 3, 'balance_date': 4, 'annual_rate': 5, 'monthly_minimum': 6})
        goals = self.rows('Savings Goals', 6, self.book['Savings Goals'].max_row - 1,
                          {'name': 1, 'target': 2, 'target_date': 3, 'opening_balance': 4, 'monthly_plan': 5})
        annual = self.rows('Annual Expenses', 6, self.book['Annual Expenses'].max_row - 1,
                           {'name': 1, 'annual_amount': 2, 'due_month': 3, 'opening_allocation': 4})
        # Recipient input region starts A6 and ends before the report bands. Do not
        # interpret Jan/Feb headers or report labels as additional recipients.
        recipients = self.rows('Sent Abroad', 6, 13, {'name': 1, 'country_note': 2, 'currency': 3})
        targets = {}
        for kind, rows in [('debt_payment', debts), ('savings_deposit', goals), ('annual_fund_payment', annual), ('remittance', recipients)]:
            names = defaultdict(list)
            for r in rows:
                if meaningful(r['raw']['name']):
                    names[r['raw']['name']].append(r['source_id'])
                else:
                    self.issue(r, 'missing_entity_name', 'name')
            targets[kind] = names
            for ids in names.values():
                if len(ids) > 1:
                    for r in rows:
                        if r['source_id'] in ids:
                            self.issue(r, 'duplicate_entity_name', 'name')
        for d in debts:
            self.amount(d, 'balance', 'reference_balance_ore')
            self.amount(d, 'monthly_minimum', 'monthly_minimum_ore')
            self.day(d, 'balance_date', 'reference_date')
            try:
                rate = decimal_value(d['raw']['annual_rate'])
                d['annual_rate_decimal'] = str(rate) if rate is not None else None
                if rate is None or rate < 0 or rate > 1:
                    self.issue(d, 'missing_or_out_of_range_annual_rate', 'annual_rate')
            except (ValueError, InvalidOperation):
                d['annual_rate_decimal'] = None
                self.issue(d, 'invalid_annual_rate', 'annual_rate')
        for g in goals:
            for field, target in [('target', 'target_ore'), ('opening_balance', 'opening_balance_ore'), ('monthly_plan', 'monthly_plan_ore')]:
                self.amount(g, field, target)
            self.day(g, 'target_date', 'target_date', required=False)
            g['opening_balance_effective_date'] = None
            if g.get('opening_balance_ore'):
                self.issue(g, 'opening_balance_date_not_in_source', 'opening_balance')
        for a in annual:
            self.amount(a, 'annual_amount', 'annual_amount_ore')
            self.amount(a, 'opening_allocation', 'opening_allocation_ore')
            a['opening_allocation_effective_date'] = None
            a['source_opening_convention'] = 'before_January_of_budget_year'

        rates = []
        for r in range(9, 19):
            if not self.entered('Settings', r, [1, 2]):
                continue
            rate = self.row('Settings', r, {'currency': 1, 'nok_per_unit': 2, 'note': 3})
            code = rate['raw']['currency']
            if r == 9 and code == '=B4':
                code = currency
                rate['resolved_from'] = 'Settings!B4'
            rate['currency'] = code
            rate['confirmed_for_conversion'] = code == 'NOK' and rate['raw']['nok_per_unit'] == 1
            if not rate['confirmed_for_conversion']:
                self.issue(rate, 'source_fx_rate_not_confirmed', 'nok_per_unit', 'warning')
            rates.append(rate)

        tx = self.rows('Transactions', 5, self.book['Transactions'].max_row,
                       {'date': 1, 'category': 2, 'description': 3, 'amount': 4, 'currency': 5, 'linked_to': 6, 'explicit_fx_rate': 7})
        tx_keys = defaultdict(list)
        for t in tx:
            self.day(t, 'date', 'effective_date')
            self.amount(t, 'amount', 'original_amount_minor_if_2dp', allow_negative=True)
            t['original_amount_decimal'] = str(t['raw']['amount']) if t['original_amount_minor_if_2dp'] is not None else None
            # Empty currency explicitly means BaseCur in this workbook.
            t['currency'] = t['raw']['currency'] or currency
            t['currency_provenance'] = 'explicit' if t['raw']['currency'] else 'workbook_blank_means_base_currency'
            cat = by_name.get(t['raw']['category'], [])
            t['category_id'] = cat[0]['source_id'] if len(cat) == 1 else None
            t['source_accounting_class'] = cat[0]['accounting_class'] if len(cat) == 1 else None
            if len(cat) != 1:
                self.issue(t, 'category_not_uniquely_resolved', 'category')
            target_type = t['source_accounting_class']
            linked = t['raw']['linked_to']
            matches = targets.get(target_type, {}).get(linked, [])
            t['linked_entity_id'] = matches[0] if len(matches) == 1 else None
            if target_type in targets and len(matches) != 1:
                self.issue(t, 'required_link_not_uniquely_resolved', 'linked_to')
            if target_type == 'expense' and meaningful(linked):
                possible = {k: v[linked] for k, v in targets.items() if linked in v}
                t['source_link_candidates'] = possible
                self.issue(t, 'expense_has_module_link_classification_review', 'linked_to')
            t['amount_nok_ore'] = None
            if t['currency'] == 'NOK':
                t['amount_nok_ore'] = t['original_amount_minor_if_2dp']
                t['fx_rate_decimal'] = '1'
            else:
                try:
                    fx = decimal_value(t['raw']['explicit_fx_rate'])
                    if fx is None or fx <= 0:
                        self.issue(t, 'missing_confirmed_transaction_fx_rate', 'explicit_fx_rate')
                    elif t['original_amount_decimal'] is not None:
                        t['fx_rate_decimal'] = str(fx)
                        t['amount_nok_ore'] = int((Decimal(t['original_amount_decimal']) * fx * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
                except (ValueError, InvalidOperation):
                    self.issue(t, 'invalid_transaction_fx_rate', 'explicit_fx_rate')
            if (t.get('amount_nok_ore') or 0) < 0:
                t['direction'] = 'refund_or_withdrawal_from_source_negative_amount'
                t['original_post_id'] = None
                self.issue(t, 'negative_amount_requires_refund_or_withdrawal_mapping', 'amount')
            else:
                t['direction'] = 'source_actual_transaction'
            t['created_by_user_id'] = None
            t['attributed_member_user_id'] = None
            t['legacy_income_person_label'] = cat[0]['raw']['note'] if len(cat) == 1 and target_type == 'income' else None
            signature = json.dumps({k: t['raw'][k] for k in t['raw']}, sort_keys=True, ensure_ascii=False)
            tx_keys[signature].append(t)
        for records in tx_keys.values():
            if len(records) > 1:
                for t in records:
                    self.issue(t, 'possible_duplicate_source_transaction_preserved', severity='warning')

        # Flags belong to source records; no source row is silently omitted.
        all_records = categories + debts + goals + annual + recipients + rates + tx
        blocking_ids = {i['record_id'] for i in self.issues if i['severity'] == 'review'}
        for r in all_records:
            r['review_state'] = 'needs_review' if r['source_id'] in blocking_ids else 'source_preserved'
        month_totals = defaultdict(lambda: defaultdict(int))
        for t in tx:
            if t['effective_date'] and t['amount_nok_ore'] is not None and t['source_accounting_class']:
                month_totals[t['effective_date'][:7]][t['source_accounting_class']] += t['amount_nok_ore']
        selected = settings['raw']['selected_month']
        month_names = ['January','February','March','April','May','June','July','August','September','October','November','December']
        selected_month = f"{settings['raw']['budget_year']}-{month_names.index(selected)+1:02d}" if selected in month_names else None
        return {'schema_version': 1, 'import_series': IMPORT_ID,
                'source': {'path': str(self.source.relative_to(ROOT)), 'sha256': digest(self.source), 'all_values_are_user_data': True},
                'target': {'scope': 'single_verified_household_only', 'household_id': None, 'owner_user_id': None,
                           'owner_binding_file': 'private-data/family-budget/profile-target.json',
                           'require_verified_owner_email': True, 'global_seed_for_new_families': False},
                'state': 'staged_not_applied_to_application', 'app_profile_populated': False,
                'policies': {'source_instructions_executed': False, 'formulas_executed': False,
                             'missing_inputs_guessed': False, 'possible_duplicates_removed': False,
                             'privacy': 'private-data excluded from Git and preview hosting',
                             'creator_assignment': 'import attribution only after owner verifies email; never infer original editor',
                             'repeat_import_key': ['household_id','import_series','source_id'],
                             'changed_source_policy': 'diff and reconcile same source IDs; never append duplicates or silently replace accepted entries',
                             'incomplete_rows': 'retain as needs-review imported drafts, visible only to target household; never manufacture confirmed totals',
                             'demo_policy': 'no fictional household/name/amount fallback in owner profile'},
                'settings': settings, 'source_selected_month': selected_month,
                'categories': categories, 'transactions': tx, 'debts': debts, 'savings_goals': goals,
                'annual_funds': annual, 'recipients': recipients, 'fx_rate_settings': rates,
                'coverage': dict(self.coverage), 'non_record_controls': self.controls,
                'issues': self.issues,
                'source_ledger_totals_ore': dict(month_totals),
                'source_totals_limitations': 'Source-date/source-kind totals only; undated rows excluded and preserved separately. Not reconciled against recalculated Excel or a live app.'}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=SOURCE)
    args = parser.parse_args()
    if not args.source.is_file():
        raise SystemExit('Private workbook not found; no source data was guessed.')
    before = digest(args.source)
    ex = Extractor(args.source)
    package = ex.extract()
    snap = ex.snapshot()
    counts = {k: len(package[k]) for k in ['categories','transactions','debts','savings_goals','annual_funds','recipients','fx_rate_settings']}
    review_rows = [r['source']['row'] for r in package['transactions'] if r['review_state'] == 'needs_review']
    validation = {'schema_version': 1, 'source_sha256': before, 'counts': counts,
                  'transaction_source_rows': package['coverage']['Transactions'],
                  'undated_transaction_rows': [r['source']['row'] for r in package['transactions'] if r['effective_date'] is None],
                  'transaction_rows_needing_review': review_rows,
                  'issue_counts': dict(Counter(i['code'] for i in ex.issues)),
                  'source_unchanged_after_read': digest(args.source) == before,
                  'formulas_recalculated': False, 'database_import_executed': False,
                  'live_profile_data_verified': False, 'all_source_input_rows_staged': True,
                  **staging_eligibility(ex.issues)}
    if not validation['source_unchanged_after_read']:
        raise SystemExit('Source changed during extraction; no package published.')
    save(PRIVATE / 'source' / 'source-receipt.json', {
        'schema_version': 1, 'original_path': 'E:/Family_Budget_Expandable.xlsx',
        'project_copy': str(args.source.relative_to(ROOT)), 'sha256': before,
        'bytes': args.source.stat().st_size, 'preserved_at': datetime.now(timezone.utc).isoformat(),
        'purpose': 'Actual family data, not a fictional template or global seed'})
    save(PRIVATE / 'import' / 'source-cell-snapshot.json', snap)
    save(PRIVATE / 'import' / 'import-package.json', package)
    save(PRIVATE / 'import' / 'validation.json', validation)
    target = PRIVATE / 'profile-target.json'
    if not target.exists():
        save(target, {'schema_version': 1, 'import_series': IMPORT_ID, 'source_sha256': before,
                      'owner_email': None, 'owner_email_verified': False, 'household_id': None,
                      'status': 'awaiting_explicit_owner_email_and_authenticated_profile',
                      'send_email_automatically': False, 'apply_to_other_families': False})
    print(json.dumps({'source_copy_verified': validation['source_unchanged_after_read'],
                      'staged_records': counts, 'transactions_requiring_review': len(review_rows),
                      'database_import_executed': False, 'private_output': 'private-data/family-budget/import/'}, ensure_ascii=False))

if __name__ == '__main__':
    main()
